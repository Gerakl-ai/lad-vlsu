"""Read a dated VLSU workbook without changing the source file."""

import argparse
import hashlib
import json
import re
from datetime import date
from pathlib import Path

import openpyxl

CLOCK = r"(?:[01]?\d|2[0-3])[.:][0-5]\d"
INTERVAL = re.compile(rf"({CLOCK})\s*[-\u2013\u2014]\s*({CLOCK})")
DAY_PREFIXES = {"пн": 0, "понедельник": 0, "вт": 1, "вторник": 1,
                "ср": 2, "среда": 2, "чт": 3, "четверг": 3,
                "пт": 4, "пятница": 4, "сб": 5, "суббота": 5,
                "вс": 6, "воскресенье": 6}


def clock(value):
    hours, minutes = re.split(r"[.:]", value)
    return f"{int(hours):02}:{minutes}"


def intervals(text):
    result = [(clock(match[1]), clock(match[2])) for match in INTERVAL.finditer(text)]
    if any(start >= end for start, end in result):
        raise ValueError(f"Invalid time interval: {text}")
    return result


def calendar_date(day, month, academic_year):
    month = int(month)
    year = academic_year[0] if month >= 8 else academic_year[1]
    return date(year, month, int(day))


def period(text):
    years = re.search(r"(20\d{2})\s*-\s*(20\d{2})", text)
    dates = re.findall(r"(?<!\d)(\d{1,2})\.(\d{2})(?!\d)", text)
    if not years or len(dates) != 2:
        raise ValueError("The source must declare a complete academic year and period")
    academic_year = tuple(map(int, years.groups()))
    if academic_year[1] != academic_year[0] + 1:
        raise ValueError("Invalid academic year")
    start, end = [calendar_date(*parts, academic_year) for parts in dates]
    if start > end:
        raise ValueError("Invalid validity period")
    return start, end, academic_year


def dated_label(text, academic_year):
    dates = re.findall(r"(?<!\d)(\d{1,2})\.(\d{2})(?!\d)", text)
    prefix = re.match(r"\s*([а-я]+)", text.lower())
    if len(dates) != 1 or not prefix or prefix[1] not in DAY_PREFIXES:
        raise ValueError(f"Ambiguous dated label: {text}")
    result = calendar_date(*dates[0], academic_year)
    if result.weekday() != DAY_PREFIXES[prefix[1]]:
        raise ValueError(f"Weekday disagrees with the source date: {text}")
    return result


def lesson_text(raw):
    # Excel sometimes splits a surname, or its initials, across wrapped lines.
    raw = re.sub(r"([А-ЯЁ][а-яё]{2,})\s*\n\s*([а-яё]{1,4})\s+([А-ЯЁ]\s*\.\s*[А-ЯЁ]\s*\.)", r"\1\2 \3", raw)
    text = re.sub(r"\s+", " ", INTERVAL.sub("", raw)).strip()
    teachers = list(re.finditer(r"[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?\s+[А-ЯЁ]\s*\.\s*[А-ЯЁ]\s*\.", text))
    rooms = list(re.finditer(r"(?<![\w-])(?:\d{1,4}|[А-ЯЁ])-\d{1,2}(?![\w-])", text))
    if len(teachers) != 1 or len(rooms) != 1:
        raise ValueError("Subject, room or teacher requires a manual source review")
    teacher = re.sub(r"\s*\.\s*", ".", teachers[0][0])
    room = rooms[0][0]
    kinds = re.findall(r"(?<!\w)(лк|лек|пр|лб)\.?(?!\w)", text, flags=re.I)
    kinds = list(dict.fromkeys("лк" if value.lower() == "лек" else value.lower() for value in kinds))
    subject = text
    for match in sorted([teachers[0], rooms[0]], key=lambda value: value.start(), reverse=True):
        subject = subject[:match.start()] + subject[match.end():]
    subject = re.sub(r"(?<!\w)(лк|лек|пр|лб|доц|проф)\.?(?!\w)", "", subject, flags=re.I)
    subject = re.sub(r"\s+", " ", subject).strip(" ,;")
    if not subject:
        raise ValueError("Empty subject after separating metadata")
    return f"{room}, {'/'.join(kinds)}, {teacher}, {subject}"


def stage(source, sheet_name, group_name, source_url, header_row=18, time_row=19, period_row=12):
    workbook = openpyxl.load_workbook(source, data_only=True)
    sheet = workbook[sheet_name]
    group_pattern = re.compile(r"(?<![\w-])" + re.escape(group_name) + r"(?![\w-])")
    headers = [cell for cell in sheet[header_row] if cell.value and group_pattern.search(str(cell.value))]
    if not headers:
        raise ValueError(f"Group is not present in the sheet: {group_name}")
    period_cells = [cell for cell in sheet[period_row] if cell.value and re.search(r"20\d{2}\s*-\s*20\d{2}", str(cell.value))]
    if len(period_cells) != 1:
        raise ValueError("Ambiguous document period")
    period_cell = period_cells[0]
    valid_from, valid_through, academic_year = period(str(period_cell.value))
    header_periods = [re.findall(r"(\d{1,2})\.(\d{2})", str(cell.value)) for cell in headers]
    if any(header_periods):
        if any(parts != header_periods[0] for parts in header_periods) or len(header_periods[0]) != 2:
            raise ValueError("Group column periods disagree")
        start, end = [calendar_date(*parts, academic_year) for parts in header_periods[0]]
        if start < valid_from or end > valid_through or start > end:
            raise ValueError("Group period is outside the document")
        valid_from, valid_through = start, end

    merges = list(sheet.merged_cells.ranges)
    merge_at = {(row, column): merge for merge in merges
                for row in range(merge.min_row, merge.max_row + 1)
                for column in range(merge.min_col, merge.max_col + 1)}

    def bounds(cell):
        merge = merge_at.get((cell.row, cell.column))
        return merge.bounds if merge else (cell.column, cell.row, cell.column, cell.row)

    def cell_at(row, column):
        merge = merge_at.get((row, column))
        return sheet.cell(merge.min_row, merge.min_col) if merge else sheet.cell(row, column)

    entries = []
    seen = set()
    errors = []
    for header in headers:
        time_headers = [cell for cell in sheet[time_row] if cell.column < header.column
                        and str(cell.value or "").strip().lower() == "часы"]
        if not time_headers:
            raise ValueError(f"Time column is absent for {header.coordinate}")
        time_column = time_headers[-1].column
        date_column = time_column - 1
        for row in sheet.iter_rows(min_row=time_row + 1):
            for cell in row:
                if cell.value is None or cell.coordinate in seen:
                    continue
                left, top, right, bottom = bounds(cell)
                if not left <= header.column <= right:
                    continue
                seen.add(cell.coordinate)
                try:
                    date_cell = cell_at(top, date_column)
                    lesson_date = dated_label(str(date_cell.value or ""), academic_year)
                    end_date_cell = cell_at(bottom, date_column)
                    if dated_label(str(end_date_cell.value or ""), academic_year) != lesson_date:
                        raise ValueError("A lesson spans more than one date")
                    if not valid_from <= lesson_date <= valid_through:
                        raise ValueError("Lesson is outside the declared group period")
                    raw = str(cell.value)
                    own_times = intervals(raw)
                    if len(own_times) > 1:
                        raise ValueError("More than one interval inside a lesson cell")
                    time_cells = {cell_at(r, time_column).coordinate: cell_at(r, time_column)
                                  for r in range(top, bottom + 1)}
                    shared_times = [item for time_cell in time_cells.values()
                                    for item in intervals(str(time_cell.value or ""))]
                    times = own_times or shared_times
                    if not times:
                        raise ValueError("No explicit time interval for this lesson")
                    name = lesson_text(raw)
                    for start, end in sorted(set(times)):
                        entries.append({"date": lesson_date.isoformat(), "start": start, "end": end, "name": name,
                                        "evidence": {"lessonCell": cell.coordinate, "rawText": raw,
                                                     "dateCell": date_cell.coordinate, "dateText": str(date_cell.value),
                                                     "headerCell": header.coordinate,
                                                     "timeCells": {key: str(value.value or "") for key, value in time_cells.items()},
                                                     "printedInterval": [start, end],
                                                     "timeInsideLesson": bool(own_times),
                                                     "bounds": [left, top, right, bottom]}})
                except ValueError as error:
                    errors.append({"cell": cell.coordinate, "message": str(error)})
    entries.sort(key=lambda entry: (entry["date"], entry["start"], entry["name"]))
    if not entries:
        raise ValueError("No lessons were extracted")
    return {"schemaVersion": 1, "purpose": "dated-spreadsheet-staging", "groupName": group_name,
            "sheet": sheet_name, "source": {"url": source_url, "sha256": hashlib.sha256(Path(source).read_bytes()).hexdigest(),
                                           "periodCell": period_cell.coordinate, "periodText": str(period_cell.value),
                                           "validFrom": valid_from.isoformat(), "validThrough": valid_through.isoformat()},
            "headers": {cell.coordinate: str(cell.value) for cell in headers}, "lessons": entries, "errors": errors}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--group", required=True)
    parser.add_argument("--source-url", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    result = stage(args.source, args.sheet, args.group, args.source_url)
    Path(args.out).write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"group": result["groupName"], "lessons": len(result["lessons"]), "errors": result["errors"]}, ensure_ascii=False))
    if result["errors"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
