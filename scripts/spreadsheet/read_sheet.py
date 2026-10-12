"""Read values, merges and visibility from XLS/XLSX without saving a workbook."""

from dataclasses import dataclass
from pathlib import Path
from types import SimpleNamespace

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange


@dataclass(frozen=True)
class SourceCell:
    row: int
    column: int
    value: object

    @property
    def coordinate(self):
        return f"{get_column_letter(self.column)}{self.row}"


class SourceSheet:
    def __init__(self, cells, rows, columns, merges, hidden_rows, hidden_columns):
        self.cells = cells
        self.rows = rows
        self.columns = columns
        self.merged_cells = SimpleNamespace(ranges=merges)
        self.hidden_rows = hidden_rows
        self.hidden_columns = hidden_columns

    def cell(self, row, column):
        return self.cells.get((row, column), SourceCell(row, column, None))

    def __getitem__(self, row):
        return [self.cell(row, column) for column in range(1, self.columns + 1)]

    def iter_rows(self, min_row=1):
        return (self[row] for row in range(min_row, self.rows + 1))


def read_sheet(source, sheet_name):
    if Path(source).suffix.lower() == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(str(source), formatting_info=True)
        sheet = workbook.sheet_by_name(sheet_name)
        cells = {(r + 1, c + 1): SourceCell(r + 1, c + 1, sheet.cell_value(r, c))
                 for r in range(sheet.nrows) for c in range(sheet.ncols) if sheet.cell_value(r, c) != ""}
        merges = [CellRange(min_col=c1 + 1, max_col=c2, min_row=r1 + 1, max_row=r2)
                  for r1, r2, c1, c2 in sheet.merged_cells]
        return SourceSheet(cells, sheet.nrows, sheet.ncols, merges,
                           {r + 1 for r, info in sheet.rowinfo_map.items() if info.hidden},
                           {c + 1 for c, info in sheet.colinfo_map.items() if info.hidden})
    workbook = openpyxl.load_workbook(source, data_only=True)
    sheet = workbook[sheet_name]
    cells = {(cell.row, cell.column): SourceCell(cell.row, cell.column, cell.value)
             for row in sheet for cell in row if cell.value is not None}
    hidden_columns = set()
    for info in sheet.column_dimensions.values():
        if info.hidden:
            hidden_columns.update(range(info.min, info.max + 1))
    result = SourceSheet(cells, sheet.max_row, sheet.max_column, list(sheet.merged_cells.ranges),
                         {r for r, info in sheet.row_dimensions.items() if info.hidden}, hidden_columns)
    workbook.close()
    return result
